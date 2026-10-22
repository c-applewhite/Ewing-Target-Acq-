"""Tests for the Ewing-branch + radius search iteration.

Covers the branch, radius, run-result, dashboard, and classification guarantees.

The suite binds to an isolated, throwaway SQLite database via APP_DB_PATH so
it never touches the production data/app.db.
"""

from __future__ import annotations

import csv
import io
import json
import os
import tempfile
import unittest
from datetime import datetime
from unittest.mock import patch

from openpyxl import load_workbook

# Isolate the database BEFORE importing the app.
_TMP_DB_FD, _TMP_DB_PATH = tempfile.mkstemp(
    suffix=".db",
    prefix="atf_ewing_test_",
)
os.close(_TMP_DB_FD)
os.environ["APP_DB_PATH"] = _TMP_DB_PATH

import app as app_module  # noqa: E402
from app import (  # noqa: E402
    app,
    db,
    classify_and_score,
    generate_coverage_points,
)
from models.database import (  # noqa: E402
    AnalystReview,
    Company,
    Exclusion,
    Location,
    Market,
    RawResult,
    RunCompanyResult,
    SearchRun,
    SourceEvidence,
)
from modules.reviewer import VALID_CLASSIFICATIONS  # noqa: E402


# Dallas / Fort Worth reference coordinates used across fixtures.
BRANCH_LAT = 32.7767
BRANCH_LNG = -96.7970


def _reset_database() -> None:
    db.drop_all()
    db.create_all()


def _make_store(
    store_id: str,
    active: bool = True,
    **overrides,
) -> Market:
    """Create an official JSON-backed Ewing location fixture."""
    official_store_id = overrides.get(
        "official_store_id",
        store_id.removeprefix("ewing-"),
    )

    branch_name = overrides.get(
        "branch_name",
        overrides.get(
            "city",
            "Dallas",
        ),
    )

    store_type = overrides.get(
        "store_type",
        "Branch",
    )

    config = {
        "source": overrides.get(
            "source",
            app_module.OFFICIAL_EWING_SOURCE,
        ),
        "active": active,
        "store_id": str(
            official_store_id
        ),
        "branch_name": str(
            branch_name
        ),
        "store_type": str(
            store_type
        ),
        "address": overrides.get(
            "address",
            "100 Main St",
        ),
        "city": overrides.get(
            "city",
            branch_name,
        ),
        "state": overrides.get(
            "state",
            "TX",
        ),
        "zip_code": overrides.get(
            "zip_code",
            "75201",
        ),
        "last_synced_at": (
            datetime.utcnow().isoformat()
        ),
    }

    store = Market(
        id=store_id,
        name=overrides.get(
            "name",
            (
                f"{official_store_id} "
                f"— {branch_name}"
            ),
        ),
        state=config["state"],
        center_lat=overrides.get(
            "lat",
            BRANCH_LAT,
        ),
        center_lng=overrides.get(
            "lng",
            BRANCH_LNG,
        ),
        primary_radius_miles=25,
        max_radius_miles=75,
        config_json=json.dumps(
            config
        ),
    )

    db.session.add(store)
    db.session.commit()

    return store


class EwingSearchTests(unittest.TestCase):
    def setUp(self) -> None:
        self.ctx = app.app_context()
        self.ctx.push()
        _reset_database()
        self.client = app.test_client()

    def tearDown(self) -> None:
        db.session.remove()
        self.ctx.pop()

    # ------------------------------------------------------------------
    # 1. /api/stores
    # ------------------------------------------------------------------
    def test_api_stores_returns_active_excludes_inactive(
        self,
    ):
        active = _make_store(
            "ewing-active-001",
            active=True,
            city="Dallas",
            branch_name="Dallas",
        )

        _make_store(
            "ewing-inactive-002",
            active=False,
            city="Austin",
            branch_name="Austin",
        )

        # Legacy Google-backed records remain available only
        # for existing run-history relationships.
        _make_store(
            "ewing-legacy-003",
            active=True,
            city="Fort Worth",
            branch_name="Fort Worth",
            source="google_places",
        )

        # A non-Ewing market must never appear in the store list.
        db.session.add(
            Market(
                id="dfw",
                name="Dallas Market",
                state="TX",
                center_lat=BRANCH_LAT,
                center_lng=BRANCH_LNG,
                config_json=json.dumps(
                    {
                        "source": (
                            app_module
                            .OFFICIAL_EWING_SOURCE
                        ),
                        "active": True,
                        "store_id": "dfw",
                    }
                ),
            )
        )

        db.session.commit()

        response = self.client.get(
            "/api/stores"
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        payload = response.get_json()

        self.assertEqual(
            len(payload),
            1,
        )

        returned_store = payload[0]

        self.assertEqual(
            returned_store["id"],
            active.id,
        )

        self.assertEqual(
            returned_store["store_id"],
            "active-001",
        )

        self.assertEqual(
            returned_store["branch_name"],
            "Dallas",
        )

        self.assertEqual(
            returned_store["store_type"],
            "Branch",
        )

        self.assertEqual(
            returned_store["label"],
            "active-001 — Dallas — Branch",
        )

        self.assertEqual(
            returned_store["source"],
            app_module.OFFICIAL_EWING_SOURCE,
        )

        self.assertIn(
            "active-001",
            returned_store["search_text"],
        )

        self.assertIn(
            "Dallas",
            returned_store["search_text"],
        )

        ids = {
            store["id"]
            for store in payload
        }

        self.assertNotIn(
            "ewing-inactive-002",
            ids,
        )

        self.assertNotIn(
            "ewing-legacy-003",
            ids,
        )

        self.assertNotIn(
            "dfw",
            ids,
        )

    def test_api_stores_survives_invalid_config_json(
        self,
    ):
        store = _make_store(
            "ewing-good-001"
        )

        bad = Market(
            id="ewing-bad-002",
            name="Ewing — Broken",
            state="TX",
            center_lat=BRANCH_LAT,
            center_lng=BRANCH_LNG,
            config_json="{not valid json",
        )

        db.session.add(bad)
        db.session.commit()

        response = self.client.get(
            "/api/stores"
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        ids = {
            store_data["id"]
            for store_data
            in response.get_json()
        }

        # The valid official store is returned.
        self.assertIn(
            store.id,
            ids,
        )

        # Invalid config is safely treated as non-official.
        self.assertNotIn(
            "ewing-bad-002",
            ids,
        )

    def test_api_runs_includes_branch_radius_and_timestamps(self):
        store = _make_store(
            "ewing-history-001",
            name="Ewing — Dallas, TX",
        )
        started_at = datetime(2026, 7, 22, 9, 15)
        completed_at = datetime(2026, 7, 22, 9, 45)
        run = SearchRun(
            market_id=store.id,
            status="completed",
            radius_miles=50,
            started_at=started_at,
            completed_at=completed_at,
        )
        db.session.add(run)
        db.session.commit()

        response = self.client.get("/api/runs")

        self.assertEqual(response.status_code, 200)
        payload = response.get_json()
        self.assertEqual(len(payload), 1)
        self.assertEqual(payload[0]["branch_name"], store.name)
        self.assertEqual(payload[0]["radius_miles"], 50)
        self.assertEqual(payload[0]["status"], "completed")
        self.assertEqual(
            payload[0]["started_at"],
            started_at.isoformat(),
        )
        self.assertEqual(
            payload[0]["completed_at"],
            completed_at.isoformat(),
        )

    # ------------------------------------------------------------------
    # 2 & 3. Coverage points
    # ------------------------------------------------------------------
    def test_small_radius_single_coverage_point(self):
        points = generate_coverage_points(
            BRANCH_LAT,
            BRANCH_LNG,
            25,
        )

        self.assertEqual(
            len(points),
            1,
        )

        self.assertAlmostEqual(
            points[0]["lat"],
            BRANCH_LAT,
        )

        self.assertAlmostEqual(
            points[0]["lng"],
            BRANCH_LNG,
        )

    def test_large_radius_multiple_coverage_points(self):
        fifty = generate_coverage_points(
            BRANCH_LAT,
            BRANCH_LNG,
            50,
        )

        seventy_five = generate_coverage_points(
            BRANCH_LAT,
            BRANCH_LNG,
            75,
        )

        self.assertGreater(
            len(fifty),
            1,
        )

        self.assertGreater(
            len(seventy_five),
            1,
        )

        # Points are deterministic across calls.
        self.assertEqual(
            fifty,
            generate_coverage_points(
                BRANCH_LAT,
                BRANCH_LNG,
                50,
            ),
        )

    # ------------------------------------------------------------------
    # 4. Out-of-radius rejection
    # ------------------------------------------------------------------
    def test_pipeline_rejects_results_outside_radius(self):
        store = _make_store(
            "ewing-radius-001"
        )

        run = SearchRun(
            market_id=store.id,
            status="pending",
            search_terms_json=json.dumps(
                ["landscape supply"]
            ),
            radius_miles=25,
        )

        db.session.add(run)
        db.session.commit()

        run_id = run.id

        near_place = {
            "id": "near-place",
            "displayName": {
                "text": "Near Yard"
            },
            "formattedAddress": (
                "1 Close Rd, Dallas, TX"
            ),
            "location": {
                "latitude": (
                    BRANCH_LAT + 0.05
                ),
                "longitude": (
                    BRANCH_LNG + 0.05
                ),
            },
        }

        far_place = {
            "id": "far-place",
            "displayName": {
                "text": "Far Yard"
            },
            "formattedAddress": (
                "1 Distant Rd, Waco, TX"
            ),
            "location": {
                # Roughly 90 miles south, beyond 25 miles.
                "latitude": (
                    BRANCH_LAT - 1.3
                ),
                "longitude": BRANCH_LNG,
            },
        }

        class _FakeCollector:
            errors: list = []

            def search(
                self,
                text_query,
                lat,
                lng,
                radius_meters,
            ):
                return [
                    near_place,
                    far_place,
                ]

            def parse_place(
                self,
                place,
            ):
                location = place.get(
                    "location",
                    {},
                )

                return {
                    "place_id": place.get(
                        "id"
                    ),
                    "name": (
                        place.get(
                            "displayName",
                            {},
                        ).get("text")
                    ),
                    "formatted_address": (
                        place.get(
                            "formattedAddress"
                        )
                    ),
                    "lat": location.get(
                        "latitude"
                    ),
                    "lng": location.get(
                        "longitude"
                    ),
                }

        import modules.collector as collector_module

        saved = (
            collector_module
            .PlacesCollector
        )

        collector_module.PlacesCollector = (
            lambda *args, **kwargs:
            _FakeCollector()
        )

        try:
            app_module.run_search_pipeline(
                run_id,
                store.id,
                ["landscape supply"],
                25,
            )
        finally:
            collector_module.PlacesCollector = (
                saved
            )

        raw = RawResult.query.filter_by(
            run_id=run_id
        ).all()

        place_ids = {
            result.place_id
            for result in raw
        }

        self.assertIn(
            "near-place",
            place_ids,
        )

        self.assertNotIn(
            "far-place",
            place_ids,
        )

    def test_runs_start_uses_defaults_when_terms_missing_or_empty(
        self,
    ):
        store = _make_store(
            "ewing-default-terms-001"
        )

        expected_terms = [
            "landscape supply",
            "hardscape supply",
            "rock and stone supplier",
            "sand and gravel supplier",
            "soil and mulch supplier",
            "paver supplier",
            "masonry supply",
        ]

        request_bodies = [
            {
                "store_id": store.id,
                "radius_miles": 25,
                "force": True,
            },
            {
                "store_id": store.id,
                "radius_miles": 25,
                "search_terms": [],
                "force": True,
            },
        ]

        with patch.object(
            app_module.threading.Thread,
            "start",
            return_value=None,
        ):
            for body in request_bodies:
                response = self.client.post(
                    "/api/runs/start",
                    json=body,
                )

                self.assertEqual(
                    response.status_code,
                    200,
                )

                run = db.session.get(
                    SearchRun,
                    response.get_json()[
                        "run_id"
                    ],
                )

                self.assertEqual(
                    json.loads(
                        run.search_terms_json
                    ),
                    expected_terms,
                )

    def test_pipeline_saves_one_raw_result_per_normalized_term(
        self,
    ):
        store = _make_store(
            "ewing-raw-dedup-001"
        )

        run = SearchRun(
            market_id=store.id,
            status="pending",
            search_terms_json=json.dumps(
                ["Landscape Supply"]
            ),
            radius_miles=50,
        )

        db.session.add(run)
        db.session.commit()

        class _FakeCollector:
            errors = []

            def search(
                self,
                **kwargs,
            ):
                return [
                    {
                        "id": (
                            "duplicate-place"
                        ),
                        "displayName": {
                            "text": (
                                "Duplicate Supply"
                            )
                        },
                        "formattedAddress": (
                            "100 Main St"
                        ),
                        "location": {
                            "latitude": (
                                BRANCH_LAT
                            ),
                            "longitude": (
                                BRANCH_LNG
                            ),
                        },
                    }
                ]

            def parse_place(
                self,
                place,
            ):
                location = place[
                    "location"
                ]

                return {
                    "place_id": (
                        place["id"]
                    ),
                    "name": (
                        place[
                            "displayName"
                        ]["text"]
                    ),
                    "formatted_address": (
                        place[
                            "formattedAddress"
                        ]
                    ),
                    "lat": (
                        location["latitude"]
                    ),
                    "lng": (
                        location["longitude"]
                    ),
                }

        class _FakeDeduplicator:
            def __init__(
                self,
                *args,
                **kwargs,
            ):
                pass

            def run(
                self,
                run_id,
            ):
                return {}

        import modules.collector as collector_module
        import modules.deduplicator as deduplicator_module

        with (
            patch.object(
                collector_module,
                "PlacesCollector",
                _FakeCollector,
            ),
            patch.object(
                deduplicator_module,
                "Deduplicator",
                _FakeDeduplicator,
            ),
            patch.object(
                app_module,
                "classify_and_score",
                return_value=0,
            ),
        ):
            app_module.run_search_pipeline(
                run.id,
                store.id,
                [
                    " Landscape Supply ",
                    "landscape supply",
                ],
                50,
            )

        raw_results = RawResult.query.filter_by(
            run_id=run.id
        ).all()

        self.assertEqual(
            len(raw_results),
            1,
        )

        self.assertEqual(
            raw_results[
                0
            ].search_term.casefold(),
            "landscape supply",
        )

        db.session.expire_all()

        self.assertEqual(
            db.session.get(
                SearchRun,
                run.id,
            ).total_raw_records,
            1,
        )

    # ------------------------------------------------------------------
    # 5. Duplicate detection key
    # ------------------------------------------------------------------
    def test_duplicate_detection_includes_store_radius_terms(
        self,
    ):
        store = _make_store(
            "ewing-dup-001"
        )

        body = {
            "store_id": store.id,
            "radius_miles": 25,
            "search_terms": [
                "landscape supply"
            ],
        }

        existing = SearchRun(
            market_id=store.id,
            status="completed",
            search_terms_json=json.dumps(
                [
                    "landscape supply"
                ]
            ),
            radius_miles=25,
        )

        db.session.add(existing)
        db.session.commit()

        duplicate = self.client.post(
            "/api/runs/start",
            json=body,
        )

        self.assertEqual(
            duplicate.status_code,
            409,
        )

        with patch.object(
            app_module.threading.Thread,
            "start",
            return_value=None,
        ):
            different_radius = (
                self.client.post(
                    "/api/runs/start",
                    json={
                        **body,
                        "radius_miles": 50,
                    },
                )
            )

            self.assertEqual(
                different_radius.status_code,
                200,
            )

            different_terms = (
                self.client.post(
                    "/api/runs/start",
                    json={
                        **body,
                        "search_terms": [
                            "stone yard"
                        ],
                    },
                )
            )

            self.assertEqual(
                different_terms.status_code,
                200,
            )

    def test_runs_start_rejects_bad_radius_and_unknown_store(
        self,
    ):
        store = _make_store(
            "ewing-val-001"
        )

        bad_radius = self.client.post(
            "/api/runs/start",
            json={
                "store_id": store.id,
                "radius_miles": 33,
                "search_terms": [
                    "stone yard"
                ],
            },
        )

        self.assertEqual(
            bad_radius.status_code,
            400,
        )

        self.assertIn(
            "10, 25, 50 or 75",
            bad_radius.get_json()[
                "error"
            ],
        )

        unknown = self.client.post(
            "/api/runs/start",
            json={
                "store_id": (
                    "ewing-nope"
                ),
                "radius_miles": 25,
                "search_terms": [
                    "stone yard"
                ],
            },
        )

        self.assertEqual(
            unknown.status_code,
            400,
        )

        self.assertEqual(
            unknown.get_json()[
                "error"
            ],
            "unknown store",
        )

    def test_runs_start_rejects_legacy_google_store(
        self,
    ):
        legacy = _make_store(
            "ewing-legacy-001",
            source="google_places",
        )

        response = self.client.post(
            "/api/runs/start",
            json={
                "store_id": legacy.id,
                "radius_miles": 25,
                "search_terms": [
                    "stone yard"
                ],
            },
        )

        self.assertEqual(
            response.status_code,
            400,
        )

        self.assertEqual(
            response.get_json()[
                "error"
            ],
            (
                "store is not an official "
                "JSON-backed Ewing location"
            ),
        )

    # ------------------------------------------------------------------
    # 6. Run-filtered results
    # ------------------------------------------------------------------
    def test_results_filtered_by_run_id(
        self,
    ):
        store = _make_store(
            "ewing-run-001"
        )

        company_a = Company(
            normalized_name="a supply",
            original_name="A Supply",
            classification=(
                "probable_target"
            ),
            overall_score=5,
        )

        company_b = Company(
            normalized_name="b supply",
            original_name="B Supply",
            classification=(
                "probable_target"
            ),
            overall_score=99,
        )

        company_c = Company(
            normalized_name="c supply",
            original_name="C Supply",
            classification=(
                "probable_target"
            ),
            overall_score=100,
        )

        db.session.add_all(
            [
                company_a,
                company_b,
                company_c,
            ]
        )

        db.session.flush()

        loc_a = Location(
            company_id=company_a.id,
            place_id="place-a",
            city="Dallas",
            state="TX",
            lat=BRANCH_LAT,
            lng=BRANCH_LNG,
        )

        loc_b = Location(
            company_id=company_b.id,
            place_id="place-b",
            city="Dallas",
            state="TX",
            lat=BRANCH_LAT,
            lng=BRANCH_LNG,
        )

        loc_c = Location(
            company_id=company_c.id,
            place_id="place-c",
            city="Dallas",
            state="TX",
            lat=BRANCH_LAT,
            lng=BRANCH_LNG,
        )

        db.session.add_all(
            [
                loc_a,
                loc_b,
                loc_c,
            ]
        )

        run = SearchRun(
            market_id=store.id,
            status="completed",
            search_terms_json=json.dumps(
                ["stone yard"]
            ),
            radius_miles=25,
        )

        other_run = SearchRun(
            market_id=store.id,
            status="completed",
            search_terms_json=json.dumps(
                ["landscape supply"]
            ),
            radius_miles=50,
        )

        db.session.add_all(
            [
                run,
                other_run,
            ]
        )

        db.session.flush()

        raw_a = RawResult(
            run_id=run.id,
            place_id="place-a",
            name="A Supply",
            search_term="stone yard",
            search_location="Dallas",
        )

        raw_b = RawResult(
            run_id=run.id,
            place_id="place-b",
            name="B Supply",
            search_term=(
                "landscape supply"
            ),
            search_location="Dallas",
        )

        raw_c = RawResult(
            run_id=other_run.id,
            place_id="place-c",
            name="C Supply",
            search_term=(
                "landscape supply"
            ),
            search_location="Dallas",
        )

        db.session.add_all(
            [
                raw_a,
                raw_b,
                raw_c,
            ]
        )

        db.session.flush()

        db.session.add_all(
            [
                SourceEvidence(
                    location_id=loc_a.id,
                    raw_result_id=raw_a.id,
                    evidence_type=(
                        "search_term"
                    ),
                    evidence_value=(
                        "stone yard"
                    ),
                ),
                SourceEvidence(
                    location_id=loc_b.id,
                    raw_result_id=raw_b.id,
                    evidence_type=(
                        "search_term"
                    ),
                    evidence_value=(
                        "landscape supply"
                    ),
                ),
                SourceEvidence(
                    location_id=loc_c.id,
                    raw_result_id=raw_c.id,
                    evidence_type=(
                        "search_term"
                    ),
                    evidence_value=(
                        "landscape supply"
                    ),
                ),
                RunCompanyResult(
                    run_id=run.id,
                    company_id=company_a.id,
                    nearest_distance_miles=3.5,
                    total_score=88,
                    score_band=(
                        "Priority Target"
                    ),
                ),
                RunCompanyResult(
                    run_id=other_run.id,
                    company_id=company_a.id,
                    nearest_distance_miles=18,
                    total_score=42,
                    score_band=(
                        "Low Priority"
                    ),
                ),
                RunCompanyResult(
                    run_id=run.id,
                    company_id=company_b.id,
                    nearest_distance_miles=6,
                    total_score=70,
                    score_band=(
                        "Strong Candidate"
                    ),
                ),
                RunCompanyResult(
                    run_id=other_run.id,
                    company_id=company_c.id,
                    nearest_distance_miles=4,
                    total_score=95,
                    score_band=(
                        "Priority Target"
                    ),
                ),
            ]
        )

        db.session.commit()

        response = self.client.get(
            (
                "/api/results?"
                f"run_id={run.id}"
            )
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        rows = response.get_json()[
            "results"
        ]

        self.assertEqual(
            [
                row["name"]
                for row in rows
            ],
            [
                "A Supply",
                "B Supply",
            ],
        )

        self.assertEqual(
            rows[0]["overall_score"],
            88,
        )

        self.assertEqual(
            rows[0][
                "distance_from_center_miles"
            ],
            3.5,
        )

        self.assertEqual(
            rows[0]["score_band"],
            "Priority Target",
        )

        self.assertEqual(
            rows[0]["evidence"],
            [
                {
                    "term": "stone yard",
                    "location": (
                        "Dallas"
                    ),
                }
            ],
        )

    def test_dashboard_stats_are_run_specific(
        self,
    ):
        store = _make_store(
            "ewing-dashboard-001"
        )

        other_store = _make_store(
            "ewing-dashboard-002",
            city="Phoenix",
            state="AZ",
            lat=33.4484,
            lng=-112.0740,
        )

        company_a = Company(
            normalized_name=(
                "alpha supply"
            ),
            original_name=(
                "Alpha Supply"
            ),
            classification=(
                "confirmed_target"
            ),
        )

        company_b = Company(
            normalized_name=(
                "beta materials"
            ),
            original_name=(
                "Beta Materials"
            ),
            classification=(
                "probable_target"
            ),
        )

        company_c = Company(
            normalized_name=(
                "other run supply"
            ),
            original_name=(
                "Other Run Supply"
            ),
            classification=(
                "confirmed_target"
            ),
        )

        db.session.add_all(
            [
                company_a,
                company_b,
                company_c,
            ]
        )

        db.session.flush()

        location_a = Location(
            company_id=company_a.id,
            place_id="dashboard-a",
            city="Dallas",
            state="TX",
            lat=BRANCH_LAT,
            lng=BRANCH_LNG,
        )

        location_b = Location(
            company_id=company_b.id,
            place_id="dashboard-b",
            city="Fort Worth",
            state="TX",
            lat=32.7555,
            lng=-97.3308,
        )

        location_c = Location(
            company_id=company_c.id,
            place_id="dashboard-c",
            city="Phoenix",
            state="AZ",
            lat=33.4484,
            lng=-112.0740,
        )

        db.session.add_all(
            [
                location_a,
                location_b,
                location_c,
            ]
        )

        run = SearchRun(
            market_id=store.id,
            status="completed",
            search_terms_json=json.dumps(
                ["landscape supply"]
            ),
            radius_miles=25,
        )

        other_run = SearchRun(
            market_id=other_store.id,
            status="completed",
            search_terms_json=json.dumps(
                [
                    "rock and stone "
                    "supplier"
                ]
            ),
            radius_miles=50,
        )

        db.session.add_all(
            [
                run,
                other_run,
            ]
        )

        db.session.flush()

        raw_a = RawResult(
            run_id=run.id,
            place_id=location_a.place_id,
            name=company_a.original_name,
            search_term=(
                "landscape supply"
            ),
            search_location="Dallas",
        )

        raw_b = RawResult(
            run_id=run.id,
            place_id=location_b.place_id,
            name=company_b.original_name,
            search_term=(
                "landscape supply"
            ),
            search_location="Dallas",
        )

        raw_c = RawResult(
            run_id=other_run.id,
            place_id=location_c.place_id,
            name=company_c.original_name,
            search_term=(
                "rock and stone supplier"
            ),
            search_location="Phoenix",
        )

        db.session.add_all(
            [
                raw_a,
                raw_b,
                raw_c,
            ]
        )

        db.session.flush()

        db.session.add_all(
            [
                SourceEvidence(
                    location_id=(
                        location_a.id
                    ),
                    raw_result_id=raw_a.id,
                ),
                SourceEvidence(
                    location_id=(
                        location_b.id
                    ),
                    raw_result_id=raw_b.id,
                ),
                SourceEvidence(
                    location_id=(
                        location_c.id
                    ),
                    raw_result_id=raw_c.id,
                ),
                RunCompanyResult(
                    run_id=run.id,
                    company_id=company_a.id,
                    nearest_distance_miles=2,
                    total_score=91,
                    score_band=(
                        "Priority Target"
                    ),
                ),
                RunCompanyResult(
                    run_id=run.id,
                    company_id=company_b.id,
                    nearest_distance_miles=20,
                    total_score=68,
                    score_band=(
                        "Strong Candidate"
                    ),
                ),
                RunCompanyResult(
                    run_id=other_run.id,
                    company_id=company_c.id,
                    nearest_distance_miles=3,
                    total_score=99,
                    score_band=(
                        "Priority Target"
                    ),
                ),
            ]
        )

        db.session.commit()

        response = self.client.get(
            (
                "/api/dashboard/stats?"
                f"run_id={run.id}"
            )
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        payload = response.get_json()

        self.assertEqual(
            payload["run"]["id"],
            run.id,
        )

        self.assertEqual(
            payload["run"]["market_id"],
            store.id,
        )

        self.assertEqual(
            payload["run"]["radius_miles"],
            25,
        )

        self.assertEqual(
            payload["raw_records"],
            2,
        )

        self.assertEqual(
            payload["unique_locations"],
            2,
        )

        self.assertEqual(
            payload["unique_companies"],
            2,
        )

        self.assertEqual(
            payload[
                "confirmed_targets"
            ],
            1,
        )

        self.assertEqual(
            payload[
                "probable_targets"
            ],
            1,
        )

        self.assertEqual(
            payload["excluded"],
            0,
        )

        self.assertEqual(
            payload["by_city"],
            {
                "Dallas": 1,
                "Fort Worth": 1,
            },
        )

        self.assertEqual(
            payload["by_band"],
            {
                "Priority Target": 1,
                "Strong Candidate": 1,
            },
        )

        self.assertEqual(
            [
                row["name"]
                for row
                in payload[
                    "top_companies"
                ]
            ],
            [
                "Alpha Supply",
                "Beta Materials",
            ],
        )

        self.assertEqual(
            payload[
                "top_companies"
            ][0]["overall_score"],
            91,
        )

    def test_dashboard_map_uses_run_branch_radius_and_locations(
        self,
    ):
        store = _make_store(
            "ewing-map-001"
        )

        other_store = _make_store(
            "ewing-map-002",
            city="Phoenix",
            state="AZ",
            lat=33.4484,
            lng=-112.0740,
        )

        company = Company(
            normalized_name=(
                "map supply"
            ),
            original_name=(
                "Map Supply"
            ),
            classification=(
                "probable_target"
            ),
        )

        other_company = Company(
            normalized_name=(
                "other map supply"
            ),
            original_name=(
                "Other Map Supply"
            ),
            classification=(
                "confirmed_target"
            ),
        )

        db.session.add_all(
            [
                company,
                other_company,
            ]
        )

        db.session.flush()

        first_location = Location(
            company_id=company.id,
            place_id="map-a",
            city="Dallas",
            state="TX",
            lat=BRANCH_LAT,
            lng=BRANCH_LNG,
        )

        second_location = Location(
            company_id=company.id,
            place_id="map-b",
            city="Irving",
            state="TX",
            lat=32.8140,
            lng=-96.9489,
        )

        other_location = Location(
            company_id=other_company.id,
            place_id="map-c",
            city="Phoenix",
            state="AZ",
            lat=33.4484,
            lng=-112.0740,
        )

        db.session.add_all(
            [
                first_location,
                second_location,
                other_location,
            ]
        )

        run = SearchRun(
            market_id=store.id,
            status="completed",
            radius_miles=25,
        )

        other_run = SearchRun(
            market_id=other_store.id,
            status="completed",
            radius_miles=50,
        )

        db.session.add_all(
            [
                run,
                other_run,
            ]
        )

        db.session.flush()

        raw_a = RawResult(
            run_id=run.id,
            place_id=(
                first_location.place_id
            ),
            name=company.original_name,
            search_term=(
                "landscape supply"
            ),
            search_location="Dallas",
        )

        raw_b = RawResult(
            run_id=run.id,
            place_id=(
                second_location.place_id
            ),
            name=company.original_name,
            search_term=(
                "landscape supply"
            ),
            search_location="Dallas",
        )

        raw_c = RawResult(
            run_id=other_run.id,
            place_id=(
                other_location.place_id
            ),
            name=(
                other_company.original_name
            ),
            search_term=(
                "rock and stone supplier"
            ),
            search_location="Phoenix",
        )

        db.session.add_all(
            [
                raw_a,
                raw_b,
                raw_c,
            ]
        )

        db.session.flush()

        db.session.add_all(
            [
                SourceEvidence(
                    location_id=(
                        first_location.id
                    ),
                    raw_result_id=raw_a.id,
                ),
                SourceEvidence(
                    location_id=(
                        second_location.id
                    ),
                    raw_result_id=raw_b.id,
                ),
                SourceEvidence(
                    location_id=(
                        other_location.id
                    ),
                    raw_result_id=raw_c.id,
                ),
                RunCompanyResult(
                    run_id=run.id,
                    company_id=company.id,
                    nearest_distance_miles=2,
                    total_score=84,
                    score_band=(
                        "Priority Target"
                    ),
                ),
                RunCompanyResult(
                    run_id=other_run.id,
                    company_id=(
                        other_company.id
                    ),
                    nearest_distance_miles=1,
                    total_score=99,
                    score_band=(
                        "Priority Target"
                    ),
                ),
            ]
        )

        db.session.commit()

        response = self.client.get(
            (
                "/api/dashboard/map-data?"
                f"run_id={run.id}"
            )
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        payload = response.get_json()

        self.assertEqual(
            payload["run_id"],
            run.id,
        )

        self.assertEqual(
            payload["market_id"],
            store.id,
        )

        self.assertEqual(
            payload["radius_miles"],
            25,
        )

        self.assertEqual(
            payload["center"],
            {
                "lat": BRANCH_LAT,
                "lng": BRANCH_LNG,
            },
        )

        self.assertEqual(
            len(
                payload[
                    "search_locations"
                ]
            ),
            1,
        )

        self.assertEqual(
            payload[
                "search_locations"
            ][0]["name"],
            store.name,
        )

        self.assertEqual(
            {
                marker["location_id"]
                for marker
                in payload["markers"]
            },
            {
                first_location.id,
                second_location.id,
            },
        )

        self.assertTrue(
            all(
                (
                    marker[
                        "company_id"
                    ] == company.id
                    and
                    marker["score"] == 84
                )
                for marker
                in payload["markers"]
            )
        )

    def test_dashboard_endpoints_reject_unknown_run(
        self,
    ):
        stats = self.client.get(
            (
                "/api/dashboard/stats?"
                "run_id=missing"
            )
        )

        map_data = self.client.get(
            (
                "/api/dashboard/map-data?"
                "run_id=missing"
            )
        )

        self.assertEqual(
            stats.status_code,
            404,
        )

        self.assertEqual(
            map_data.status_code,
            404,
        )

    # ------------------------------------------------------------------
    # 7. Run-specific exports
    # ------------------------------------------------------------------
    def test_excel_and_csv_exports_are_run_specific(
        self,
    ):
        store = _make_store(
            "ewing-export-001"
        )

        other_store = _make_store(
            "ewing-export-002",
            city="Phoenix",
            state="AZ",
            lat=33.4484,
            lng=-112.0740,
        )

        ranked = Company(
            normalized_name=(
                "ranked supply"
            ),
            original_name=(
                "Ranked Supply"
            ),
            classification=(
                "confirmed_target"
            ),
            overall_score=1,
            analyst_notes=(
                "Keep this note"
            ),
            ownership_type=(
                "independent"
            ),
            is_reviewed=True,
        )

        review = Company(
            normalized_name=(
                "review supply"
            ),
            original_name=(
                "Review Supply"
            ),
            classification=(
                "manual_review"
            ),
        )

        excluded = Company(
            normalized_name=(
                "excluded supply"
            ),
            original_name=(
                "Excluded Supply"
            ),
            classification="excluded",
        )

        other = Company(
            normalized_name=(
                "other run supply"
            ),
            original_name=(
                "Other Run Supply"
            ),
            classification=(
                "confirmed_target"
            ),
            overall_score=100,
        )

        db.session.add_all(
            [
                ranked,
                review,
                excluded,
                other,
            ]
        )

        db.session.flush()

        ranked_location = Location(
            company_id=ranked.id,
            place_id="export-ranked",
            raw_address="1 Dallas Rd",
            city="Dallas",
            state="TX",
            lat=BRANCH_LAT,
            lng=BRANCH_LNG,
            rating=4.7,
            review_count=120,
        )

        ranked_other_location = Location(
            company_id=ranked.id,
            place_id=(
                "export-ranked-other"
            ),
            raw_address=(
                "1 Phoenix Rd"
            ),
            city="Phoenix",
            state="AZ",
            lat=33.4484,
            lng=-112.0740,
        )

        review_location = Location(
            company_id=review.id,
            place_id="export-review",
            city="Fort Worth",
            state="TX",
            lat=32.7555,
            lng=-97.3308,
        )

        excluded_location = Location(
            company_id=excluded.id,
            place_id=(
                "export-excluded"
            ),
            city="Irving",
            state="TX",
            lat=32.8140,
            lng=-96.9489,
        )

        other_location = Location(
            company_id=other.id,
            place_id="export-other",
            city="Phoenix",
            state="AZ",
            lat=33.4484,
            lng=-112.0740,
        )

        db.session.add_all(
            [
                ranked_location,
                ranked_other_location,
                review_location,
                excluded_location,
                other_location,
            ]
        )

        run = SearchRun(
            market_id=store.id,
            status="completed",
            search_terms_json=json.dumps(
                ["landscape supply"]
            ),
            radius_miles=25,
        )

        other_run = SearchRun(
            market_id=other_store.id,
            status="completed",
            search_terms_json=json.dumps(
                ["masonry supply"]
            ),
            radius_miles=50,
        )

        db.session.add_all(
            [
                run,
                other_run,
            ]
        )

        db.session.flush()

        raw_rows = [
            RawResult(
                run_id=run.id,
                place_id=(
                    ranked_location
                    .place_id
                ),
                name=ranked.original_name,
                search_term=(
                    "landscape supply"
                ),
                search_location=(
                    "Dallas"
                ),
            ),
            RawResult(
                run_id=run.id,
                place_id=(
                    review_location
                    .place_id
                ),
                name=review.original_name,
                search_term=(
                    "landscape supply"
                ),
                search_location=(
                    "Dallas"
                ),
            ),
            RawResult(
                run_id=run.id,
                place_id=(
                    excluded_location
                    .place_id
                ),
                name=(
                    excluded.original_name
                ),
                search_term=(
                    "landscape supply"
                ),
                search_location=(
                    "Dallas"
                ),
            ),
            RawResult(
                run_id=other_run.id,
                place_id=(
                    ranked_other_location
                    .place_id
                ),
                name=ranked.original_name,
                search_term=(
                    "masonry supply"
                ),
                search_location=(
                    "Phoenix"
                ),
            ),
            RawResult(
                run_id=other_run.id,
                place_id=(
                    other_location
                    .place_id
                ),
                name=other.original_name,
                search_term=(
                    "masonry supply"
                ),
                search_location=(
                    "Phoenix"
                ),
            ),
        ]

        db.session.add_all(
            raw_rows
        )

        db.session.flush()

        db.session.add_all(
            [
                SourceEvidence(
                    location_id=(
                        ranked_location.id
                    ),
                    raw_result_id=(
                        raw_rows[0].id
                    ),
                ),
                SourceEvidence(
                    location_id=(
                        review_location.id
                    ),
                    raw_result_id=(
                        raw_rows[1].id
                    ),
                ),
                SourceEvidence(
                    location_id=(
                        excluded_location.id
                    ),
                    raw_result_id=(
                        raw_rows[2].id
                    ),
                ),
                SourceEvidence(
                    location_id=(
                        ranked_other_location
                        .id
                    ),
                    raw_result_id=(
                        raw_rows[3].id
                    ),
                ),
                SourceEvidence(
                    location_id=(
                        other_location.id
                    ),
                    raw_result_id=(
                        raw_rows[4].id
                    ),
                ),
                RunCompanyResult(
                    run_id=run.id,
                    company_id=ranked.id,
                    nearest_distance_miles=(
                        3.25
                    ),
                    product_fit_score=35,
                    geo_fit_score=22,
                    business_quality_score=18,
                    data_confidence_score=13,
                    total_score=88,
                    score_band=(
                        "Priority Target"
                    ),
                ),
                RunCompanyResult(
                    run_id=run.id,
                    company_id=review.id,
                    nearest_distance_miles=20,
                    total_score=55,
                    score_band=(
                        "Monitor / Investigate"
                    ),
                ),
                RunCompanyResult(
                    run_id=run.id,
                    company_id=excluded.id,
                    nearest_distance_miles=10,
                    total_score=0,
                    score_band=(
                        "Low Priority"
                    ),
                ),
                RunCompanyResult(
                    run_id=other_run.id,
                    company_id=ranked.id,
                    nearest_distance_miles=1,
                    total_score=12,
                    score_band=(
                        "Low Priority"
                    ),
                ),
                RunCompanyResult(
                    run_id=other_run.id,
                    company_id=other.id,
                    nearest_distance_miles=2,
                    total_score=99,
                    score_band=(
                        "Priority Target"
                    ),
                ),
                Exclusion(
                    company_id=excluded.id,
                    location_id=(
                        excluded_location.id
                    ),
                    exclusion_reason=(
                        "Not an independent "
                        "supplier"
                    ),
                    exclusion_keywords_matched=(
                        "national chain"
                    ),
                ),
            ]
        )

        db.session.commit()

        csv_response = self.client.get(
            (
                "/api/export/csv?"
                f"run_id={run.id}"
            )
        )

        self.assertEqual(
            csv_response.status_code,
            200,
        )

        csv_text = (
            csv_response.get_data(
                as_text=True
            )
        )

        csv_rows = list(
            csv.DictReader(
                io.StringIO(
                    csv_text
                )
            )
        )

        self.assertEqual(
            len(csv_rows),
            1,
        )

        self.assertEqual(
            csv_rows[0][
                "Company Name"
            ],
            "Ranked Supply",
        )

        self.assertEqual(
            csv_rows[0]["Cities"],
            "Dallas",
        )

        self.assertEqual(
            csv_rows[0][
                "Location Count"
            ],
            "1",
        )

        self.assertEqual(
            csv_rows[0][
                "Total Score"
            ],
            "88.0",
        )

        self.assertEqual(
            csv_rows[0][
                "Product Fit"
            ],
            "35.0",
        )

        self.assertEqual(
            csv_rows[0][
                "Nearest Distance (mi)"
            ],
            "3.25",
        )

        self.assertEqual(
            csv_rows[0][
                "Analyst Notes"
            ],
            "Keep this note",
        )

        self.assertNotIn(
            "Other Run Supply",
            csv_text,
        )

        self.assertNotIn(
            "Phoenix",
            csv_text,
        )

        excel_response = self.client.get(
            (
                "/api/export?"
                f"run_id={run.id}"
            )
        )

        self.assertEqual(
            excel_response.status_code,
            200,
        )

        workbook = load_workbook(
            io.BytesIO(
                excel_response.data
            ),
            read_only=True,
            data_only=True,
        )

        ranked_rows = list(
            workbook[
                "Ranked Targets"
            ].iter_rows(
                values_only=True
            )
        )

        self.assertEqual(
            ranked_rows[1][1],
            "Ranked Supply",
        )

        self.assertEqual(
            ranked_rows[1][2],
            1,
        )

        self.assertEqual(
            ranked_rows[1][3],
            "Dallas",
        )

        self.assertEqual(
            ranked_rows[1][16],
            88,
        )

        review_rows = list(
            workbook[
                "Manual Review"
            ].iter_rows(
                values_only=True
            )
        )

        self.assertEqual(
            review_rows[1][0],
            "Review Supply",
        )

        excluded_rows = list(
            workbook[
                "Excluded"
            ].iter_rows(
                values_only=True
            )
        )

        self.assertEqual(
            excluded_rows[1][0],
            "Excluded Supply",
        )

        self.assertEqual(
            excluded_rows[1][3],
            (
                "Not an independent "
                "supplier"
            ),
        )

        location_rows = list(
            workbook[
                "All Locations"
            ].iter_rows(
                values_only=True
            )
        )

        exported_location_names = {
            row[0]
            for row
            in location_rows[1:]
        }

        self.assertEqual(
            exported_location_names,
            {
                "Ranked Supply",
                "Review Supply",
                "Excluded Supply",
            },
        )

        self.assertNotIn(
            "Other Run Supply",
            exported_location_names,
        )

        self.assertNotIn(
            "Phoenix",
            {
                row[3]
                for row
                in location_rows[1:]
            },
        )

        summary_rows = list(
            workbook[
                "Search Summary"
            ].iter_rows(
                values_only=True
            )
        )

        self.assertTrue(
            summary_rows[1:]
        )

        self.assertEqual(
            {
                row[0]
                for row
                in summary_rows[1:]
            },
            {
                run.id
            },
        )

        self.assertEqual(
            {
                row[1]
                for row
                in summary_rows[1:]
            },
            {
                store.name
            },
        )

        self.assertEqual(
            {
                row[2]
                for row
                in summary_rows[1:]
            },
            {
                25
            },
        )

    def test_exports_require_a_valid_run_id(
        self,
    ):
        self.assertEqual(
            self.client.get(
                "/api/export"
            ).status_code,
            400,
        )

        self.assertEqual(
            self.client.get(
                "/api/export/csv"
            ).status_code,
            400,
        )

        self.assertEqual(
            self.client.get(
                (
                    "/api/export?"
                    "run_id=missing"
                )
            ).status_code,
            404,
        )

        self.assertEqual(
            self.client.get(
                (
                    "/api/export/csv?"
                    "run_id=missing"
                )
            ).status_code,
            404,
        )

    # ------------------------------------------------------------------
    # 8. Analyst overrides survive reprocessing
    # ------------------------------------------------------------------
    def test_analyst_override_survives_reprocessing(
        self,
    ):
        store = _make_store(
            "ewing-override-001"
        )

        company = Company(
            normalized_name=(
                "override supply"
            ),
            original_name=(
                "Override Supply"
            ),
            classification=(
                "probable_target"
            ),
            overall_score=50,
        )

        db.session.add(company)
        db.session.flush()

        location = Location(
            company_id=company.id,
            place_id="place-o",
            city="Dallas",
            state="TX",
            lat=BRANCH_LAT,
            lng=BRANCH_LNG,
            google_category=(
                "landscape_supply_store"
            ),
        )

        run = SearchRun(
            market_id=store.id,
            status="completed",
            search_terms_json=json.dumps(
                ["landscape supply"]
            ),
            radius_miles=25,
        )

        db.session.add_all(
            [
                location,
                run,
            ]
        )

        db.session.flush()

        raw = RawResult(
            run_id=run.id,
            place_id="place-o",
            name="Override Supply",
            search_term=(
                "landscape supply"
            ),
            search_location="Dallas",
        )

        db.session.add(raw)
        db.session.flush()

        db.session.add_all(
            [
                SourceEvidence(
                    location_id=(
                        location.id
                    ),
                    raw_result_id=raw.id,
                    evidence_type=(
                        "search_term"
                    ),
                    evidence_value=(
                        "landscape supply"
                    ),
                ),
                AnalystReview(
                    company_id=company.id,
                    analyst_classification=(
                        "excluded"
                    ),
                    classification_reason=(
                        "Analyst decision"
                    ),
                    reviewed_at=(
                        datetime.utcnow()
                    ),
                ),
            ]
        )

        db.session.commit()

        classify_and_score(
            run.id,
            company_ids=[
                company.id
            ],
        )

        refreshed = db.session.get(
            Company,
            company.id,
        )

        run_result = (
            RunCompanyResult.query
            .filter_by(
                run_id=run.id,
                company_id=company.id,
            )
            .one()
        )

        self.assertEqual(
            refreshed.classification,
            "excluded",
        )

        self.assertEqual(
            run_result.total_score,
            0,
        )

    def test_scoring_preserves_results_for_each_run(
        self,
    ):
        store_a = _make_store(
            "ewing-store-a",
            lat=BRANCH_LAT,
            lng=BRANCH_LNG,
        )

        store_b = _make_store(
            "ewing-store-b",
            city="Phoenix",
            state="AZ",
            lat=33.4484,
            lng=-112.0740,
        )

        company = Company(
            normalized_name=(
                "shared landscape supply"
            ),
            original_name=(
                "Shared Landscape Supply"
            ),
        )

        db.session.add(company)
        db.session.flush()

        location = Location(
            company_id=company.id,
            place_id="shared-place",
            city="Dallas",
            state="TX",
            lat=BRANCH_LAT,
            lng=BRANCH_LNG,
            google_category=(
                "landscape_supply_store"
            ),
            business_status=(
                "OPERATIONAL"
            ),
        )

        run_a = SearchRun(
            market_id=store_a.id,
            radius_miles=25,
            status="completed",
        )

        run_b = SearchRun(
            market_id=store_b.id,
            radius_miles=75,
            status="completed",
        )

        db.session.add_all(
            [
                location,
                run_a,
                run_b,
            ]
        )

        db.session.flush()

        raw_a = RawResult(
            run_id=run_a.id,
            place_id="shared-place",
            name=company.original_name,
            search_term=(
                "landscape supply"
            ),
            search_location="Dallas",
        )

        raw_b = RawResult(
            run_id=run_b.id,
            place_id="shared-place",
            name=company.original_name,
            search_term=(
                "rock and stone supplier"
            ),
            search_location="Phoenix",
        )

        db.session.add_all(
            [
                raw_a,
                raw_b,
            ]
        )

        db.session.flush()

        db.session.add_all(
            [
                SourceEvidence(
                    location_id=location.id,
                    raw_result_id=raw_a.id,
                    evidence_type=(
                        "search_term"
                    ),
                    evidence_value=(
                        "landscape supply"
                    ),
                ),
                SourceEvidence(
                    location_id=location.id,
                    raw_result_id=raw_b.id,
                    evidence_type=(
                        "search_term"
                    ),
                    evidence_value=(
                        "rock and stone supplier"
                    ),
                ),
            ]
        )

        db.session.commit()

        classify_and_score(
            run_a.id,
            company_ids=[
                company.id
            ],
        )

        first_result = (
            RunCompanyResult.query
            .filter_by(
                run_id=run_a.id,
                company_id=company.id,
            )
            .one()
        )

        first_distance = (
            first_result
            .nearest_distance_miles
        )

        first_score = (
            first_result.total_score
        )

        classify_and_score(
            run_b.id,
            company_ids=[
                company.id
            ],
        )

        db.session.expire_all()

        preserved_result = (
            RunCompanyResult.query
            .filter_by(
                run_id=run_a.id,
                company_id=company.id,
            )
            .one()
        )

        second_result = (
            RunCompanyResult.query
            .filter_by(
                run_id=run_b.id,
                company_id=company.id,
            )
            .one()
        )

        refreshed_location = (
            db.session.get(
                Location,
                location.id,
            )
        )

        refreshed_company = (
            db.session.get(
                Company,
                company.id,
            )
        )

        self.assertEqual(
            (
                RunCompanyResult.query
                .filter_by(
                    company_id=company.id
                )
                .count()
            ),
            2,
        )

        self.assertEqual(
            (
                preserved_result
                .nearest_distance_miles
            ),
            first_distance,
        )

        self.assertEqual(
            preserved_result.total_score,
            first_score,
        )

        self.assertNotEqual(
            (
                second_result
                .nearest_distance_miles
            ),
            first_distance,
        )

        self.assertIsNone(
            (
                refreshed_location
                .distance_from_center_miles
            )
        )

        self.assertEqual(
            refreshed_company.overall_score,
            0,
        )

    # ------------------------------------------------------------------
    # 9. possible_target removed
    # ------------------------------------------------------------------
    def test_possible_target_not_valid(
        self,
    ):
        self.assertNotIn(
            "possible_target",
            VALID_CLASSIFICATIONS,
        )

        self.assertEqual(
            VALID_CLASSIFICATIONS,
            {
                "confirmed_target",
                "probable_target",
                "manual_review",
                "excluded",
            },
        )


if __name__ == "__main__":
    unittest.main()