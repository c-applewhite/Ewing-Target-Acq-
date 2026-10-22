"""Run-specific Excel and CSV exports for the Acquisition Target Finder."""

from __future__ import annotations

import csv
import io
import math
from collections import defaultdict
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="1F2A44")
HEADER_FONT = Font(color="FFFFFF", bold=True)

BAND_FILLS = {
    "Priority Target": PatternFill("solid", fgColor="C6EFCE"),
    "Strong Candidate": PatternFill("solid", fgColor="BDD7EE"),
    "Monitor / Investigate": PatternFill("solid", fgColor="FFEB9C"),
    "Low Priority": PatternFill("solid", fgColor="D9D9D9"),
}

PRODUCT_GROUPS = {
    "Bulk Materials": ["bulk_materials"],
    "Decorative Rock": ["decorative_rock"],
    "Sand and Gravel": ["sand_gravel"],
    "Soil, Compost, or Mulch": [
        "soil_compost_mulch",
        "soil_compost",
        "mulch",
    ],
    "Pavers or Retaining Walls": [
        "pavers_retaining_walls",
        "pavers",
        "retaining_walls",
    ],
    "Masonry or Natural Stone": [
        "masonry_natural_stone",
        "masonry",
        "natural_stone",
    ],
    "Delivery": ["delivery"],
}


class Exporter:
    """Build exports for one search run."""

    def __init__(self, db, models, run_id: str) -> None:
        self.db = db
        self.m = models
        self.run_id = str(run_id or "").strip()
        self.run = db.session.get(
            self.m.SearchRun,
            self.run_id,
        )

        if not self.run:
            raise ValueError("search run not found")

        self.market = db.session.get(
            self.m.Market,
            self.run.market_id,
        )

    # ------------------------------------------------------------------
    # Run-scoped query helpers
    # ------------------------------------------------------------------

    def _company_result_rows(
        self,
        classifications: Optional[list[str]] = None,
    ) -> list[tuple[object, object]]:
        query = (
            self.db.session.query(
                self.m.Company,
                self.m.RunCompanyResult,
            )
            .join(
                self.m.RunCompanyResult,
                self.m.RunCompanyResult.company_id
                == self.m.Company.id,
            )
            .filter(
                self.m.RunCompanyResult.run_id == self.run_id
            )
        )

        if classifications:
            query = query.filter(
                self.m.Company.classification.in_(classifications)
            )

        return query.order_by(
            self.m.RunCompanyResult.total_score.desc(),
            self.m.Company.original_name.asc(),
        ).all()

    def _locations_for_company(
        self,
        company_id: int,
    ) -> list[object]:
        return (
            self.m.Location.query
            .join(
                self.m.SourceEvidence,
                self.m.SourceEvidence.location_id
                == self.m.Location.id,
            )
            .join(
                self.m.RawResult,
                self.m.RawResult.id
                == self.m.SourceEvidence.raw_result_id,
            )
            .filter(
                self.m.RawResult.run_id == self.run_id,
                self.m.Location.company_id == company_id,
            )
            .order_by(self.m.Location.id.asc())
            .distinct()
            .all()
        )

    # ------------------------------------------------------------------
    # Location helpers
    # ------------------------------------------------------------------

    @staticmethod
    def _representative_location(locations: list) -> Optional[object]:
        if not locations:
            return None

        return sorted(
            locations,
            key=lambda location: (
                -(location.review_count or 0),
                location.id,
            ),
        )[0]

    @staticmethod
    def _first_nonempty(locations: list, attribute: str) -> Any:
        for location in sorted(locations, key=lambda item: item.id):
            value = getattr(location, attribute, None)

            if value not in (None, ""):
                return value

        return ""

    @staticmethod
    def _joined_unique(values: list[Any]) -> str:
        cleaned = {
            str(value).strip()
            for value in values
            if value is not None and str(value).strip()
        }

        return ", ".join(sorted(cleaned))

    @staticmethod
    def _haversine_miles(
        lat1: float,
        lng1: float,
        lat2: float,
        lng2: float,
    ) -> float:
        earth_radius_miles = 3958.7613
        lat1_radians = math.radians(lat1)
        lat2_radians = math.radians(lat2)
        delta_lat = math.radians(lat2 - lat1)
        delta_lng = math.radians(lng2 - lng1)

        value = (
            math.sin(delta_lat / 2) ** 2
            + math.cos(lat1_radians)
            * math.cos(lat2_radians)
            * math.sin(delta_lng / 2) ** 2
        )

        return earth_radius_miles * 2 * math.atan2(
            math.sqrt(value),
            math.sqrt(1 - value),
        )

    def _location_distance(self, location) -> Optional[float]:
        if (
            not self.market
            or self.market.center_lat is None
            or self.market.center_lng is None
            or location.lat is None
            or location.lng is None
        ):
            return None

        return round(
            self._haversine_miles(
                self.market.center_lat,
                self.market.center_lng,
                location.lat,
                location.lng,
            ),
            2,
        )

    # ------------------------------------------------------------------
    # Product and company rows
    # ------------------------------------------------------------------

    def _product_string(self, company_id: int) -> str:
        product_record = self.m.ProductCategory.query.filter_by(
            company_id=company_id
        ).first()

        if not product_record:
            return ""

        detected = [
            label
            for label, field_names in PRODUCT_GROUPS.items()
            if any(
                bool(getattr(product_record, field_name, False))
                for field_name in field_names
            )
        ]

        return ", ".join(detected)

    def _company_rows(
        self,
        classifications: list[str],
        include_rank: bool,
    ) -> list[dict[str, Any]]:
        rows: list[dict[str, Any]] = []

        for rank, (company, run_result) in enumerate(
            self._company_result_rows(classifications),
            start=1,
        ):
            locations = self._locations_for_company(company.id)
            representative = self._representative_location(locations)

            row = {
                "Company Name": company.original_name or "",
                "Location Count": len(locations),
                "Cities": self._joined_unique(
                    [location.city for location in locations]
                ),
                "States": self._joined_unique(
                    [location.state for location in locations]
                ),
                "Website": self._first_nonempty(
                    locations,
                    "website_raw",
                ),
                "Phone": self._first_nonempty(
                    locations,
                    "phone_normalized",
                ),
                "Classification": company.classification or "",
                "Product Categories": self._product_string(company.id),
                "Rating": representative.rating if representative else "",
                "Review Count": (
                    representative.review_count if representative else ""
                ),
                "Nearest Distance (mi)": (
                    round(run_result.nearest_distance_miles, 2)
                    if run_result.nearest_distance_miles is not None
                    else ""
                ),
                "Product Fit": round(
                    run_result.product_fit_score or 0,
                    1,
                ),
                "Geographic Fit": round(
                    run_result.geo_fit_score or 0,
                    1,
                ),
                "Business Quality": round(
                    run_result.business_quality_score or 0,
                    1,
                ),
                "Data Confidence": round(
                    run_result.data_confidence_score or 0,
                    1,
                ),
                "Total Score": round(run_result.total_score or 0, 1),
                "Score Band": run_result.score_band or "",
                "Ownership": company.ownership_type or "unknown",
                "Analyst Notes": company.analyst_notes or "",
                "Review Status": (
                    "Reviewed" if company.is_reviewed else "Pending"
                ),
            }

            if include_rank:
                row["Rank"] = rank

            rows.append(row)

        return rows

    # ------------------------------------------------------------------
    # Worksheet formatting
    # ------------------------------------------------------------------

    def _write_sheet(
        self,
        worksheet,
        rows: list[dict[str, Any]],
        headers: list[str],
        band_column: Optional[str] = None,
    ) -> None:
        worksheet.append(headers)

        for cell in worksheet[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        band_index = None

        if band_column and band_column in headers:
            band_index = headers.index(band_column) + 1

        for row in rows:
            worksheet.append([row.get(header, "") for header in headers])

            if band_index:
                fill = BAND_FILLS.get(row.get(band_column, ""))

                if fill:
                    worksheet.cell(
                        row=worksheet.max_row,
                        column=band_index,
                    ).fill = fill

        worksheet.freeze_panes = "A2"

        if rows:
            worksheet.auto_filter.ref = (
                f"A1:{get_column_letter(len(headers))}"
                f"{worksheet.max_row}"
            )

        for column_number, header in enumerate(headers, start=1):
            maximum_length = len(str(header))

            for row in rows:
                maximum_length = max(
                    maximum_length,
                    len(str(row.get(header, ""))),
                )

            worksheet.column_dimensions[
                get_column_letter(column_number)
            ].width = min(max(maximum_length + 2, 10), 55)

    # ------------------------------------------------------------------
    # Excel workbook
    # ------------------------------------------------------------------

    def build_workbook(self) -> io.BytesIO:
        workbook = Workbook()

        ranked_headers = [
            "Rank",
            "Company Name",
            "Location Count",
            "Cities",
            "States",
            "Website",
            "Phone",
            "Classification",
            "Product Categories",
            "Rating",
            "Review Count",
            "Nearest Distance (mi)",
            "Product Fit",
            "Geographic Fit",
            "Business Quality",
            "Data Confidence",
            "Total Score",
            "Score Band",
            "Ownership",
            "Analyst Notes",
            "Review Status",
        ]
        review_headers = [
            header for header in ranked_headers if header != "Rank"
        ]

        ranked_sheet = workbook.active
        ranked_sheet.title = "Ranked Targets"
        self._write_sheet(
            ranked_sheet,
            self._company_rows(
                ["confirmed_target", "probable_target"],
                include_rank=True,
            ),
            ranked_headers,
            band_column="Score Band",
        )

        review_sheet = workbook.create_sheet("Manual Review")
        self._write_sheet(
            review_sheet,
            self._company_rows(
                ["manual_review"],
                include_rank=False,
            ),
            review_headers,
            band_column="Score Band",
        )

        excluded_sheet = workbook.create_sheet("Excluded")
        excluded_headers = [
            "Company Name",
            "Cities",
            "States",
            "Exclusion Reason",
            "Keywords Matched",
            "Analyst Notes",
        ]
        excluded_rows: list[dict[str, Any]] = []

        for company, _run_result in self._company_result_rows(
            ["excluded"]
        ):
            locations = self._locations_for_company(company.id)
            exclusion = self.m.Exclusion.query.filter_by(
                company_id=company.id
            ).order_by(
                self.m.Exclusion.excluded_at.desc()
            ).first()

            excluded_rows.append(
                {
                    "Company Name": company.original_name or "",
                    "Cities": self._joined_unique(
                        [location.city for location in locations]
                    ),
                    "States": self._joined_unique(
                        [location.state for location in locations]
                    ),
                    "Exclusion Reason": (
                        exclusion.exclusion_reason if exclusion else ""
                    ),
                    "Keywords Matched": (
                        exclusion.exclusion_keywords_matched
                        if exclusion
                        else ""
                    ),
                    "Analyst Notes": company.analyst_notes or "",
                }
            )

        self._write_sheet(
            excluded_sheet,
            excluded_rows,
            excluded_headers,
        )

        locations_sheet = workbook.create_sheet("All Locations")
        location_headers = [
            "Company",
            "Location ID",
            "Place ID",
            "City",
            "State",
            "ZIP",
            "Address",
            "Phone",
            "Website",
            "Google Category",
            "Rating",
            "Reviews",
            "Business Status",
            "Distance (mi)",
            "Classification",
            "Company Score",
        ]
        location_rows: list[dict[str, Any]] = []

        for company, run_result in self._company_result_rows():
            for location in self._locations_for_company(company.id):
                location_rows.append(
                    {
                        "Company": company.original_name or "",
                        "Location ID": location.id,
                        "Place ID": location.place_id or "",
                        "City": location.city or "",
                        "State": location.state or "",
                        "ZIP": location.zip_code or "",
                        "Address": location.raw_address or "",
                        "Phone": location.phone_normalized or "",
                        "Website": location.website_raw or "",
                        "Google Category": location.google_category or "",
                        "Rating": location.rating,
                        "Reviews": location.review_count,
                        "Business Status": location.business_status or "",
                        "Distance (mi)": self._location_distance(location),
                        "Classification": company.classification or "",
                        "Company Score": round(
                            run_result.total_score or 0,
                            1,
                        ),
                    }
                )

        self._write_sheet(
            locations_sheet,
            location_rows,
            location_headers,
        )

        summary_sheet = workbook.create_sheet("Search Summary")
        summary_headers = [
            "Run ID",
            "Ewing Branch",
            "Radius (mi)",
            "Run Status",
            "Search Term",
            "Search Location",
            "Raw Results",
            "Unique Places",
        ]
        summary_data: dict[
            tuple[str, str],
            dict[str, Any],
        ] = defaultdict(
            lambda: {
                "raw_count": 0,
                "place_ids": set(),
            }
        )

        for raw_result in self.m.RawResult.query.filter_by(
            run_id=self.run_id
        ).all():
            key = (
                raw_result.search_term or "",
                raw_result.search_location or "",
            )
            summary_data[key]["raw_count"] += 1

            if raw_result.place_id:
                summary_data[key]["place_ids"].add(raw_result.place_id)

        summary_rows = [
            {
                "Run ID": self.run_id,
                "Ewing Branch": (
                    self.market.name if self.market else self.run.market_id
                ),
                "Radius (mi)": self.run.radius_miles,
                "Run Status": self.run.status,
                "Search Term": search_term,
                "Search Location": search_location,
                "Raw Results": values["raw_count"],
                "Unique Places": len(values["place_ids"]),
            }
            for (search_term, search_location), values
            in sorted(summary_data.items())
        ]

        self._write_sheet(
            summary_sheet,
            summary_rows,
            summary_headers,
        )

        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)
        return output

    # ------------------------------------------------------------------
    # CSV
    # ------------------------------------------------------------------

    def build_csv(self) -> str:
        rows = self._company_rows(
            ["confirmed_target", "probable_target"],
            include_rank=True,
        )
        headers = [
            "Rank",
            "Company Name",
            "Location Count",
            "Cities",
            "States",
            "Website",
            "Phone",
            "Classification",
            "Product Categories",
            "Rating",
            "Review Count",
            "Nearest Distance (mi)",
            "Product Fit",
            "Geographic Fit",
            "Business Quality",
            "Data Confidence",
            "Total Score",
            "Score Band",
            "Ownership",
            "Analyst Notes",
            "Review Status",
        ]
        output = io.StringIO()
        writer = csv.DictWriter(
            output,
            fieldnames=headers,
            extrasaction="ignore",
        )
        writer.writeheader()

        for row in rows:
            writer.writerow(row)

        return output.getvalue()
